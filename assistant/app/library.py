"""Your private library: a folder on the Pi that Jarvis can search and answer from, with the local model only.

Every file is read once, cut into short passages and indexed with a small embedding model running in
Ollama on the Pi. A question finds the closest passages; the local model answers from them. Nothing here
is ever sent to Claude (see privacy.py).

Reads: text, Markdown, CSV, JSON, HTML, PDF (with OCR for scanned pages), Word, Excel, PowerPoint,
OpenDocument text, images (text in the picture + a short description), sound files (transcribed), and
your voice notes.
"""
from __future__ import annotations

import asyncio
import base64
import csv
import html
import io
import json
import logging
import os
import re
import sqlite3
import subprocess
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np

from . import config, privacy

log = logging.getLogger("library")

TEXT_EXT = {".txt", ".md", ".markdown", ".log", ".rst", ".org", ".ini", ".cfg", ".yaml", ".yml", ".toml", ".tex"}
KINDS = {
    **{e: "text" for e in TEXT_EXT},
    ".csv": "table", ".tsv": "table", ".json": "json", ".html": "web page", ".htm": "web page",
    ".pdf": "PDF", ".docx": "Word", ".xlsx": "spreadsheet", ".xlsm": "spreadsheet", ".pptx": "slides",
    ".odt": "document", ".rtf": "document",
    ".jpg": "image", ".jpeg": "image", ".png": "image", ".webp": "image", ".gif": "image", ".bmp": "image",
    ".tif": "image", ".tiff": "image", ".heic": "image",
    ".mp3": "sound", ".wav": "sound", ".m4a": "sound", ".ogg": "sound", ".opus": "sound", ".flac": "sound",
    ".aac": "sound", ".wma": "sound",
}
OLD_FORMATS = {".doc": "Word 97", ".xls": "Excel 97", ".ppt": "PowerPoint 97"}
SKIP_EXT = {".db", ".sqlite", ".sqlite3", ".db-journal", ".db-wal", ".db-shm", ".tmp", ".part", ".lock"}
CHUNK, OVERLAP = 900, 150


@dataclass
class Hit:
    path: str
    loc: str
    text: str
    score: float

    @property
    def name(self) -> str:
        return Path(self.path).name


# =============================================================================
# Reading files into (location, text) parts
# =============================================================================
def _run(cmd: list[str], data: bytes | None = None, timeout: int = 120) -> bytes:
    return subprocess.run(cmd, input=data, capture_output=True, timeout=timeout, check=True).stdout


def ocr_image(path: Path) -> str:
    try:
        return _run(["tesseract", str(path), "stdout", "--psm", "3"], timeout=90).decode("utf-8", "replace").strip()
    except FileNotFoundError:
        log.warning("tesseract isn't installed: text in pictures won't be read")
    except Exception as exc:
        log.info("OCR failed for %s: %s", path.name, exc)
    return ""


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def read_table(path: Path) -> list[tuple[str, str]]:
    text = read_text_file(path)
    dialect = "excel-tab" if path.suffix.lower() == ".tsv" else None
    if dialect is None:
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = "excel"
    rows = list(csv.reader(io.StringIO(text), dialect))
    return table_parts(rows, "")


def table_parts(rows: list[list], where: str) -> list[tuple[str, str]]:
    """Rows as 'header: value' lines, a few rows per part, so each passage makes sense on its own."""
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    if not body:
        return [(where or "table", " | ".join(header))]
    parts, block, start = [], [], 2
    for i, r in enumerate(body, start=2):
        cells = [f"{header[j] if j < len(header) and header[j] else f'column {j + 1}'}: {v}"
                 for j, v in enumerate(r) if v]
        block.append("; ".join(cells))
        if sum(len(b) for b in block) > CHUNK - 100:
            parts.append((f"{where}rows {start}-{i}", "\n".join(block)))
            block, start = [], i + 1
    if block:
        parts.append((f"{where}rows {start}-{start + len(block) - 1}", "\n".join(block)))
    return parts


def flatten_json(value, prefix: str = "") -> list[str]:
    out: list[str] = []
    if isinstance(value, dict):
        for k, v in value.items():
            out += flatten_json(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(value, list):
        for i, v in enumerate(value):
            out += flatten_json(v, f"{prefix}[{i}]")
    else:
        out.append(f"{prefix}: {value}")
    return out


def read_json(path: Path) -> list[tuple[str, str]]:
    data = json.loads(read_text_file(path))
    if isinstance(data, list) and data and all(isinstance(x, dict) for x in data[:50]):
        keys: list[str] = []
        for row in data:
            for k in row:
                if k not in keys:
                    keys.append(k)
        rows = [keys] + [[json.dumps(r.get(k)) if isinstance(r.get(k), (dict, list)) else r.get(k) for k in keys]
                         for r in data]
        return table_parts(rows, "")
    return [("", "\n".join(flatten_json(data)))]


def read_html(path: Path) -> str:
    text = read_text_file(path)
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</h\d>|</li>|</tr>", "\n", text)
    return html.unescape(re.sub(r"<[^>]+>", " ", text))


def read_pdf(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("it's password-protected")
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            parts.append((f"page {i}", text))
    if not parts:   # a scan: read the pictures of the pages
        with tempfile.TemporaryDirectory() as tmp:
            try:
                _run(["pdftoppm", "-r", "150", "-f", "1", "-l", str(config.LIBRARY_OCR_PAGES), "-png",
                      str(path), f"{tmp}/p"], timeout=300)
            except FileNotFoundError:
                raise ValueError("it's a scan and the PDF tools aren't installed")
            for png in sorted(Path(tmp).glob("p*.png")):
                text = ocr_image(png)
                if text:
                    num = int(re.sub(r"\D", "", png.stem) or 0)
                    parts.append((f"page {num} (scanned)", text))
    return parts


def read_docx(path: Path) -> list[tuple[str, str]]:
    import docx
    d = docx.Document(str(path))
    lines, parts, heading = [], [], ""
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if p.style is not None and (p.style.name or "").lower().startswith("heading"):
            if lines:
                parts.append((heading, "\n".join(lines)))
                lines = []
            heading = t
        lines.append(t)
    if lines:
        parts.append((heading, "\n".join(lines)))
    for n, table in enumerate(d.tables, start=1):
        rows = [[c.text for c in row.cells] for row in table.rows]
        parts += table_parts(rows, f"table {n}, ")
    return parts


def read_xlsx(path: Path) -> list[tuple[str, str]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
            if len(rows) > 20000:
                break
        parts += table_parts(rows, f"sheet {ws.title}, ")
    wb.close()
    return parts


def read_pptx(path: Path) -> list[tuple[str, str]]:
    from pptx import Presentation
    parts = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if texts:
            parts.append((f"slide {i}", "\n".join(texts)))
    return parts


def read_odt(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("content.xml").decode("utf-8", "replace")
    xml = re.sub(r"</text:(p|h)>", "\n", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def read_rtf(path: Path) -> str:
    text = read_text_file(path)
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\'[0-9a-f]{2}|\\[a-z]+-?\d* ?|[{}]", "", text)
    return text


# =============================================================================
# The index
# =============================================================================
class Library:
    def __init__(self, stt=None, http=None) -> None:
        self.dir = config.LIBRARY_DIR
        self.notes_dir = config.NOTES_DIR
        self.stt = stt
        self.http = http
        self.db_path = config.DATA_DIR / "library" / "index.db"
        self._lock = threading.RLock()
        self._db: sqlite3.Connection | None = None
        self._matrix: np.ndarray | None = None
        self._ids: np.ndarray | None = None
        self._dirty = True
        self.queue: asyncio.Queue[str] = asyncio.Queue()
        self.queued: set[str] = set()
        self.busy: str | None = None
        self.embed_ready = False
        self.embed_error = ""
        self.vision_ready: bool | None = None
        self.last_scan = 0.0
        self.scan_now = asyncio.Event()

    # ------------------------------------------------------------ storage
    def db(self) -> sqlite3.Connection:
        if self._db is None:
            self.db_path.parent.mkdir(parents=True, exist_ok=True)
            db = sqlite3.connect(self.db_path, check_same_thread=False)
            db.executescript("""
                CREATE TABLE IF NOT EXISTS files (path TEXT PRIMARY KEY, mtime REAL, size INTEGER, kind TEXT,
                    status TEXT, detail TEXT, chunks INTEGER DEFAULT 0, indexed_at REAL);
                CREATE TABLE IF NOT EXISTS chunks (id INTEGER PRIMARY KEY, path TEXT, loc TEXT, text TEXT, vec BLOB);
                CREATE INDEX IF NOT EXISTS chunks_path ON chunks(path);
            """)
            try:
                db.execute("CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(text, tokenize='porter')")
                self._fts = True
            except sqlite3.OperationalError:
                self._fts = False
            db.commit()
            os.chmod(self.db_path, 0o600)
            self._db = db
        return self._db

    def stats(self) -> dict:
        with self._lock:
            db = self.db()
            rows = db.execute("SELECT status, COUNT(*), SUM(chunks) FROM files GROUP BY status").fetchall()
        out = {"indexed": 0, "waiting": 0, "failed": 0, "skipped": 0, "passages": 0}
        for status, n, c in rows:
            out[status if status in out else "failed"] = n
            out["passages"] += c or 0
        out["waiting"] += len(self.queued)
        return out

    def files(self) -> list[dict]:
        with self._lock:
            rows = self.db().execute("SELECT path, kind, size, status, detail, chunks, indexed_at FROM files "
                                     "ORDER BY path").fetchall()
        return [dict(zip(("path", "kind", "size", "status", "detail", "chunks", "indexed_at"), r)) for r in rows]

    def empty(self) -> bool:
        with self._lock:
            return self.db().execute("SELECT COUNT(*) FROM chunks").fetchone()[0] == 0

    # ------------------------------------------------------------ models (Ollama on the Pi)
    async def ensure_models(self, pull) -> None:
        """pull(name) makes sure an Ollama model is there (Brain.pull_model)."""
        try:
            await pull(config.LIBRARY_EMBED_MODEL)
            await self.embed(["warm up"], query=True)
            self.embed_ready, self.embed_error = True, ""
            log.info("Library search model %s ready", config.LIBRARY_EMBED_MODEL)
        except Exception as exc:
            self.embed_error = str(exc)
            log.error("Library search model %s isn't working: %s", config.LIBRARY_EMBED_MODEL, exc)

    async def embed(self, texts: list[str], query: bool = False) -> np.ndarray:
        prefix = "search_query: " if query else "search_document: "
        if "nomic" not in config.LIBRARY_EMBED_MODEL:
            prefix = ""
        r = await self.http.post(f"{config.OLLAMA_URL}/api/embed", json={
            "model": config.LIBRARY_EMBED_MODEL, "input": [prefix + t for t in texts], "keep_alive": "24h"},
            timeout=300)
        r.raise_for_status()
        vecs = np.asarray(r.json()["embeddings"], dtype=np.float32)
        norms = np.linalg.norm(vecs, axis=1, keepdims=True)
        return vecs / np.maximum(norms, 1e-9)

    async def describe_image(self, path: Path, pull) -> str:
        if not config.LIBRARY_VISION_MODEL:
            return ""
        if self.vision_ready is None:
            try:
                await pull(config.LIBRARY_VISION_MODEL)
                self.vision_ready = True
            except Exception as exc:
                log.warning("Picture descriptions are off (%s couldn't load: %s)", config.LIBRARY_VISION_MODEL, exc)
                self.vision_ready = False
        if not self.vision_ready:
            return ""
        try:   # small JPEG first: quicker for the model and kind to memory
            jpeg = await asyncio.to_thread(_run, ["ffmpeg", "-v", "error", "-i", str(path), "-frames:v", "1",
                                                  "-vf", "scale='min(768,iw)':-2", "-f", "image2", "-c:v", "mjpeg",
                                                  "-q:v", "4", "-"], None, 60)
            r = await self.http.post(f"{config.OLLAMA_URL}/api/generate", json={
                "model": config.LIBRARY_VISION_MODEL, "stream": False, "keep_alive": "2m",
                "prompt": "Describe this picture in one or two plain sentences: who or what is in it, where, "
                          "and any writing you can see.",
                "images": [base64.b64encode(jpeg).decode()]}, timeout=600)
            r.raise_for_status()
            return r.json().get("response", "").strip()
        except Exception as exc:
            log.info("Couldn't describe %s: %s", path.name, exc)
            return ""

    # ------------------------------------------------------------ scanning
    def walk(self) -> dict[str, tuple[float, int, Path]]:
        """Every file Jarvis should know about: {library path: (mtime, size, real path)}."""
        found: dict[str, tuple[float, int, Path]] = {}
        if self.dir.is_dir():
            for root, dirs, names in os.walk(self.dir):
                dirs[:] = [d for d in dirs if not d.startswith(".")]
                for n in names:
                    if n.startswith(".") or n.startswith("~$"):
                        continue
                    p = Path(root) / n
                    try:
                        st = p.stat()
                    except OSError:
                        continue
                    found[p.relative_to(self.dir).as_posix()] = (st.st_mtime, st.st_size, p)
        if config.LIBRARY_VOICE_NOTES and self.notes_dir.is_dir():
            for p in self.notes_dir.glob("*.txt"):
                if re.fullmatch(r"\d{14}", p.stem):
                    st = p.stat()
                    found[f"voice notes/{p.name}"] = (st.st_mtime, st.st_size, p)
        return found

    def real_path(self, path: str) -> Path:
        if path.startswith("voice notes/"):
            return self.notes_dir / path.split("/", 1)[1]
        return self.dir / path

    async def scan(self) -> None:
        found = await asyncio.to_thread(self.walk)
        with self._lock:
            known = {p: (m, s) for p, m, s in self.db().execute("SELECT path, mtime, size FROM files")}
        gone = [p for p in known if p not in found]
        for p in gone:
            self.forget(p)
        if gone:
            log.info("Library: %d file(s) removed", len(gone))
        for path, (mtime, size, _real) in sorted(found.items()):
            if known.get(path) != (mtime, size) and path not in self.queued:
                self.queued.add(path)
                self.queue.put_nowait(path)
        self.last_scan = time.time()

    async def watch(self) -> None:
        while True:
            try:
                await self.scan()
            except Exception:
                log.exception("Library scan failed")
            try:
                await asyncio.wait_for(self.scan_now.wait(), config.LIBRARY_SCAN_S)
            except asyncio.TimeoutError:
                pass
            self.scan_now.clear()

    def forget(self, path: str) -> None:
        with self._lock:
            db = self.db()
            ids = [r[0] for r in db.execute("SELECT id FROM chunks WHERE path=?", (path,))]
            db.execute("DELETE FROM chunks WHERE path=?", (path,))
            if self._fts and ids:
                db.executemany("DELETE FROM chunks_fts WHERE rowid=?", [(i,) for i in ids])
            db.execute("DELETE FROM files WHERE path=?", (path,))
            db.commit()
            self._dirty = True

    def set_file(self, path: str, mtime: float, size: int, kind: str, status: str, detail: str, chunks: int) -> None:
        with self._lock:
            self.db().execute("INSERT OR REPLACE INTO files VALUES (?,?,?,?,?,?,?,?)",
                              (path, mtime, size, kind, status, detail, chunks, time.time()))
            self.db().commit()

    # ------------------------------------------------------------ indexing
    async def run(self, pull, ready: asyncio.Event | None = None) -> None:
        asyncio.create_task(self.watch())
        if ready is not None:
            await ready.wait()   # let the main model finish loading first (Ollama does one thing at a time)
        await self.ensure_models(pull)
        while True:
            path = await self.queue.get()
            self.busy = path
            try:
                await self.index(path, pull)
            except Exception as exc:
                log.exception("Library: couldn't index %s", path)
                self.set_file(path, 0, 0, "", "failed", f"error: {exc}", 0)
            finally:
                self.queued.discard(path)
                self.busy = None

    async def read(self, path: str, real: Path, pull) -> tuple[str, list[tuple[str, str]]]:
        """(kind, [(location, text)])."""
        if path.startswith("voice notes/"):
            title = ""
            try:
                title = json.loads(real.with_suffix(".json").read_text()).get("title", "")
            except Exception:
                pass
            when = time.strftime("%d %b %Y %H:%M", time.strptime(real.stem, "%Y%m%d%H%M%S"))
            return "voice note", [(f"recorded {when}" + (f", '{title}'" if title else ""), read_text_file(real))]
        ext = real.suffix.lower()
        kind = KINDS.get(ext, "")
        if kind == "text":
            return kind, [("", read_text_file(real))]
        if kind == "table":
            return kind, await asyncio.to_thread(read_table, real)
        if kind == "json":
            return kind, await asyncio.to_thread(read_json, real)
        if kind == "web page":
            return kind, [("", await asyncio.to_thread(read_html, real))]
        if kind == "PDF":
            return kind, await asyncio.to_thread(read_pdf, real)
        if kind == "Word":
            return kind, await asyncio.to_thread(read_docx, real)
        if kind == "spreadsheet":
            return kind, await asyncio.to_thread(read_xlsx, real)
        if kind == "slides":
            return kind, await asyncio.to_thread(read_pptx, real)
        if kind == "document":
            fn = read_odt if ext == ".odt" else read_rtf
            return kind, [("", await asyncio.to_thread(fn, real))]
        if kind == "image":
            text = await asyncio.to_thread(ocr_image, real) if ext not in (".gif", ".heic") else ""
            desc = await self.describe_image(real, pull)
            parts = []
            if desc:
                parts.append(("what's in the picture", desc))
            if len(re.findall(r"[A-Za-z]{3,}", text)) >= 3:   # ignore stray OCR noise
                parts.append(("text in the picture", text))
            return kind, parts
        if kind == "sound":
            if self.stt is None:
                return kind, []
            pcm = await asyncio.to_thread(_run, ["ffmpeg", "-v", "error", "-i", str(real), "-ac", "1", "-ar", "16000",
                                                 "-f", "s16le", "-"], None, 900)
            text = await asyncio.to_thread(self.stt.transcribe_note, pcm)
            return kind, [(f"transcript ({len(pcm) // 32000 // 60}:{len(pcm) // 32000 % 60:02d} long)", text)]
        return kind, []

    async def index(self, path: str, pull) -> None:
        real = self.real_path(path)
        try:
            st = real.stat()
        except FileNotFoundError:
            self.forget(path)
            return
        ext = real.suffix.lower()
        if ext in SKIP_EXT:
            self.set_file(path, st.st_mtime, st.st_size, "database", "skipped",
                          "a database file: add it under Databases to ask questions about it", 0)
            return
        if ext in OLD_FORMATS:
            self.set_file(path, st.st_mtime, st.st_size, OLD_FORMATS[ext], "skipped",
                          f"old {OLD_FORMATS[ext]} format: save it as {ext}x to include it", 0)
            return
        if not path.startswith("voice notes/") and ext not in KINDS:
            self.set_file(path, st.st_mtime, st.st_size, ext.lstrip(".") or "file", "skipped",
                          "not a kind of file Jarvis can read", 0)
            return
        if st.st_size > config.LIBRARY_MAX_MB * 1024 * 1024:
            self.set_file(path, st.st_mtime, st.st_size, KINDS.get(ext, ""), "skipped",
                          f"bigger than {config.LIBRARY_MAX_MB:g} MB (LIBRARY_MAX_MB)", 0)
            return
        t0 = time.monotonic()
        try:
            kind, parts = await self.read(path, real, pull)
        except Exception as exc:
            reason = str(exc) if isinstance(exc, ValueError) else f"couldn't read it ({exc.__class__.__name__})"
            log.info("Library: %s: %s", path, exc)
            self.set_file(path, st.st_mtime, st.st_size, KINDS.get(ext, ""), "failed", reason, 0)
            return
        chunks = [(loc, piece) for loc, text in parts for piece in split_text(text)]
        if not chunks:
            self.forget(path)
            self.set_file(path, st.st_mtime, st.st_size, kind, "skipped", "no words found in it", 0)
            return
        if not self.embed_ready:
            await self.ensure_models(pull)
            if not self.embed_ready:
                raise RuntimeError("the search model isn't loaded yet")
        name = Path(path).name
        vecs = []
        for i in range(0, len(chunks), 16):
            batch = chunks[i:i + 16]
            vecs.append(await self.embed([f"{name} {loc}: {txt}" for loc, txt in batch]))
        vec = np.vstack(vecs).astype(np.float16)
        with self._lock:
            self.forget(path)
            db = self.db()
            for (loc, txt), v in zip(chunks, vec):
                cur = db.execute("INSERT INTO chunks (path, loc, text, vec) VALUES (?,?,?,?)",
                                 (path, loc, txt, v.tobytes()))
                if self._fts:
                    db.execute("INSERT INTO chunks_fts (rowid, text) VALUES (?,?)", (cur.lastrowid, f"{name} {txt}"))
            db.commit()
            self.set_file(path, st.st_mtime, st.st_size, kind, "indexed", "", len(chunks))
            self._dirty = True
        log.info("Library: indexed %s (%s, %d passages, %.1fs)", path, kind, len(chunks), time.monotonic() - t0)

    # ------------------------------------------------------------ searching
    def _load(self) -> None:
        with self._lock:
            if not self._dirty:
                return
            rows = self.db().execute("SELECT id, vec FROM chunks").fetchall()
            if rows:
                self._ids = np.array([r[0] for r in rows])
                self._matrix = np.vstack([np.frombuffer(r[1], dtype=np.float16) for r in rows]).astype(np.float32)
            else:
                self._ids, self._matrix = None, None
            self._dirty = False

    def _keyword(self, query: str, limit: int = 20) -> dict[int, int]:
        if not self._fts:
            return {}
        words = [w for w in re.findall(r"[a-zA-Z0-9]{3,}", query.lower()) if w not in _STOP]
        if not words:
            return {}
        q = " OR ".join(f'"{w}"' for w in words)
        with self._lock:
            try:
                rows = self.db().execute("SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? ORDER BY rank LIMIT ?",
                                         (q, limit)).fetchall()
            except sqlite3.OperationalError:
                return {}
        return {r[0]: i for i, r in enumerate(rows)}

    async def search(self, query: str, k: int = 5) -> list[Hit]:
        if not query.strip():
            return []
        await asyncio.to_thread(self._load)
        scores: dict[int, float] = {}
        if self._matrix is not None and self.embed_ready:
            try:
                q = (await self.embed([query], query=True))[0]
                sims = self._matrix @ q
                top = np.argsort(-sims)[:20]
                scores = {int(self._ids[i]): float(sims[i]) for i in top}
            except Exception as exc:
                log.warning("Library search model failed (%s): using word search only", exc)
        keyword = await asyncio.to_thread(self._keyword, query)
        for cid, rank in keyword.items():   # a word match nudges a passage up (or in, without the model)
            scores[cid] = scores.get(cid, 0.0 if self.embed_ready and self._matrix is not None else 0.45) \
                + 0.04 * (1 - rank / 20)
        if not scores:
            return []
        best = [kv for kv in sorted(scores.items(), key=lambda kv: -kv[1]) if kv[1] > 0.05][:k]
        with self._lock:
            out = []
            for cid, score in best:
                row = self.db().execute("SELECT path, loc, text FROM chunks WHERE id=?", (cid,)).fetchone()
                if row:
                    out.append(Hit(row[0], row[1], row[2], round(score, 3)))
        return out

    @staticmethod
    def context(hits: list[Hit], max_chars: int = 3600) -> str:
        """The passages for the local model, marked so they can never reach Claude."""
        lines, used = [], 0
        for i, h in enumerate(hits, start=1):
            text = re.sub(r"\s+", " ", h.text).strip()
            room = max_chars - used
            if room < 200:
                break
            text = text[:room]
            used += len(text)
            where = f" ({h.loc})" if h.loc else ""
            lines.append(f"[{i}] {h.name}{where}: {text}")
        return privacy.mark(
            "Passages from the user's own private files, found for this question:\n" + "\n".join(lines) +
            "\n\nAnswer from these passages only. Say which file it came from in a natural way (the file name "
            "without its extension, e.g. 'your insurance letter says...'). If they don't contain the answer, say "
            "you couldn't find it in their files rather than guessing. Never invent figures or dates.")


_STOP = set("the and for are was what when where which who how does did can you your with from that this have has "
            "about into there their them they then than our out not but all any its it's i'm my me".split())


def split_text(text: str) -> list[str]:
    """Passages of about CHUNK characters, overlapping a little, broken at paragraph or sentence ends."""
    text = re.sub(r"[ \t\r\f\v]+", " ", text or "")
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if len(re.findall(r"\w", text)) < 3:
        return []
    out, start = [], 0
    while start < len(text):
        end = min(len(text), start + CHUNK)
        if end < len(text):
            cut = max(text.rfind("\n\n", start + CHUNK // 2, end), text.rfind(". ", start + CHUNK // 2, end))
            if cut > start:
                end = cut + 1
        piece = text[start:end].strip()
        if piece:
            out.append(piece)
        if end >= len(text):
            break
        start = max(end - OVERLAP, start + 1)
        nxt = text.find(" ", start)   # start the next passage on a word
        if 0 <= nxt < end:
            start = nxt + 1
    return out


# =============================================================================
# "Check my files for ...": an explicit search
# =============================================================================
_FILES = r"(?:files?|documents?|docs|library|papers|folder|notes)"
EXPLICIT_RE = [
    re.compile(rf"^\W*(?:please\s+)?(?:can you\s+|could you\s+)?(?:check|search|look (?:in|through|at)|go through)\s+"
               rf"(?:my|the)\s+(?:private\s+)?{_FILES}\s*(?:for|about|on|to see|and tell me)?\s*[,:]?\s*(?P<q>.+)$", re.I),
    re.compile(rf"^\W*(?:what|which)\s+(?:do|does)\s+my\s+(?:private\s+)?{_FILES}\s+(?:say|tell me|have)\s+(?:about\s+)?(?P<q>.+)$", re.I),
    re.compile(rf"^\W*(?:according to|in|from|based on)\s+my\s+(?:private\s+)?{_FILES}\s*[,:]?\s*(?P<q>.+)$", re.I),
    re.compile(rf"^(?P<q>.+?)[,\s]+(?:in|from|according to)\s+my\s+(?:private\s+)?{_FILES}\W*$", re.I),
]


def explicit_query(text: str) -> str | None:
    for regex in EXPLICIT_RE:
        m = regex.match(text.strip())
        if m:
            q = m.group("q").strip(" ?.!,")
            if q:
                return q
    return None
